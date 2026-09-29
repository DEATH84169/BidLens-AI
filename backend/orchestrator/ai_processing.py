"""Document extraction with complete page accounting and source-linked claims.

Extraction establishes what was submitted, not whether an issuer authenticated it.
"""

import os
import re
from pathlib import Path
import pymupdf
import docx
import openpyxl
import pandas as pd

_ocr_engine = None


def get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        import onnxruntime

        onnxruntime.disable_telemetry_events()
        from rapidocr_onnxruntime import RapidOCR

        _ocr_engine = RapidOCR()
    return _ocr_engine


def ocr_text(data):
    results, _ = get_ocr_engine()(data)
    return "\n".join(r[1] for r in results or [])


def extract_pages(path):
    ext = Path(path).suffix.lower()
    pages = []

    def add(text, page, error=None):
        pages.append(
            {
                "page": page,
                "text": text,
                "error": error or (None if text.strip() else "No readable text"),
            }
        )

    if ext == ".pdf":
        with pymupdf.open(path) as document:
            if len(document) > 200:
                raise ValueError(
                    "PDF exceeds the 200-page limit; split it into attachments."
                )
            for i, page in enumerate(document):
                text = page.get_text()
                if len(text.strip()) < 30:
                    try:
                        text = ocr_text(page.get_pixmap(dpi=150).tobytes("png"))
                    except Exception:
                        add("", i + 1, "OCR failed; manual page review required")
                        continue
                add(text, i + 1)
    elif ext == ".docx":
        document = docx.Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells)
            for t in document.tables
            for row in t.rows
        )
        add(
            text, None
        )  # DOCX pagination is layout-dependent; do not invent page numbers.
    elif ext == ".xlsx":
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for sheet in workbook:
                add(
                    "\n".join(
                        " | ".join(str(v) for v in row if v is not None)
                        for row in sheet.values
                    ),
                    sheet.title,
                )
        finally:
            workbook.close()
    elif ext == ".csv":
        add(pd.read_csv(path).to_string(index=False), None)
    elif ext in {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"}:
        try:
            add(ocr_text(str(path)), 1)
        except Exception:
            add("", 1, "OCR failed; manual review required")
    else:
        raise ValueError("Unsupported format; convert legacy DOC/XLS to DOCX/XLSX.")
    return pages


def extract_document_text(path):
    pages = extract_pages(path)
    return (
        "\n".join(p["text"] for p in pages),
        len(pages),
        Path(path).suffix[1:].upper(),
    )


def money(value, unit):
    amount = float(value.replace(",", ""))
    return amount * (
        10000000
        if unit and unit.lower().startswith("cr")
        else 100000
        if unit and unit.lower().startswith("lakh")
        else 1
    )


PATTERNS = {
    "gstin": r"\b\d{2}[A-Z]{5}\d{4}[A-Z][A-Z\d]Z[A-Z\d]\b",
    "pan": r"\b[A-Z]{5}\d{4}[A-Z]\b",
    "udyam": r"\bUDYAM-[A-Z]{2}-\d{2}-\d{7}\b",
}


def extract_from_pages(pages, filename):
    fields, evidence = {}, {}

    def record(key, value, page, match):
        evidence.setdefault(key, []).append(
            {
                "filename": filename,
                "page": page["page"],
                "text": match.group(0),
                "value": value,
            }
        )
        fields.setdefault(key, value)

    for page in pages:
        text = page["text"]
        for key, pattern in PATTERNS.items():
            for m in re.finditer(pattern, text, re.I):
                record(key, m.group(0).upper(), page, m)
        for key, pattern in {
            "vendor_name": r"(?:Bidder Name|Company Name|Vendor Name|Name of (?:the )?Bidder)\s*:\s*([^\n\r]+)",
            "cin": r"\b([LU]\d{5}[A-Z]{2}\d{4}[A-Z]{3}\d{6})\b",
            "nic2008": r"NIC\s*2008\s*(?:Code)?\s*:\s*(\d{2,5})\b",
        }.items():
            for m in re.finditer(pattern, text, re.I):
                record(key, m.group(1).strip(), page, m)
        for m in re.finditer(
            r"(?:Annual\s+Turnover|Turnover)\s*[:=]?\s*(?:INR|Rs\.?|₹)?\s*([\d,]+(?:\.\d+)?)\s*(crores?|cr\b|lakhs?)",
            text,
            re.I,
        ):
            record("turnover_cr", money(m.group(1), m.group(2)) / 10000000, page, m)
        for m in re.finditer(
            r"(?:Total\s+(?:Quoted\s+)?(?:Price|Quote|Amount)|Quoted\s+Price)\s*[:=]\s*(?:INR|Rs\.?|₹)?\s*([\d,]+(?:\.\d+)?)\s*(crores?|cr\b|lakhs?)?",
            text,
            re.I,
        ):
            record("total_quote_inr", money(m.group(1), m.group(2)), page, m)
        for m in re.finditer(
            r"(?:Local\s+Content|Local\s+Value\s+Addition)\s*[:=]?\s*(\d+(?:\.\d+)?)\s*%",
            text,
            re.I,
        ):
            val = float(m.group(1))
            if 0 <= val <= 100:
                record("local_content_pct", val, page, m)
        # Only explicit labelled commitments, not mentions anywhere in boilerplate.
        for m in re.finditer(
            r"(?:Warranty\s*(?:Offered|Period)?\s*[:=]\s*)(\d+)\s*[- ]?\s*(years?|months?)",
            text,
            re.I,
        ):
            value = float(m.group(1)) / (
                12 if m.group(2).lower().startswith("month") else 1
            )
            record("warranty_years", value, page, m)
        for m in re.finditer(
            r"EMD\s+(?:Amount\s+)?(?:Submitted|Paid|Deposited)\s*[:=]\s*(?:INR|Rs\.?|₹)?\s*([\d,]+(?:\.\d+)?)",
            text,
            re.I,
        ):
            record("emd_amount_inr", money(m.group(1), None), page, m)
        # A claimed exemption is never a verified exemption.
        for m in re.finditer(
            r"(?:MSME\s+(?:Exemption\s+)?Claimed|MSE\s+Exemption\s+Claimed)\s*:\s*yes\b",
            text,
            re.I,
        ):
            record("msme_claimed", True, page, m)
    fields.update(
        filename=filename,
        file_type=Path(filename).suffix[1:].upper(),
        vendor_name=fields.get("vendor_name", "Unidentified bidder"),
        page_count=len(pages),
        evidence=evidence,
        unread_pages=[
            {"filename": filename, "page": p["page"], "reason": p["error"]}
            for p in pages
            if p["error"]
        ],
        raw_text="\n".join(p["text"] for p in pages),
    )
    fields["all_pans"] = list(
        dict.fromkeys(e["value"] for e in evidence.get("pan", []))
    )
    fields["all_gstins"] = list(
        dict.fromkeys(e["value"] for e in evidence.get("gstin", []))
    )
    fields["msme_claimed"] = bool(fields.get("msme_claimed") or fields.get("udyam"))
    fields["is_msme"] = False  # Cannot establish legal status through document text.
    fields["warranty"] = (
        f"{fields['warranty_years']:g}-Year"
        if fields.get("warranty_years") is not None
        else "Not extracted"
    )
    fields["emd_status"] = (
        "CLAIMED_SUBMITTED"
        if fields.get("emd_amount_inr") is not None
        else "NOT_VERIFIED"
    )
    return fields


def extract_document_data(path):
    try:
        pages = extract_pages(path)
    except Exception as exc:
        raise ValueError(
            "Attachment could not be parsed; review file format and readability."
        ) from exc
    return extract_from_pages(pages, Path(path).name)


def merge_documents(documents):
    if not documents:
        raise ValueError("At least one attachment required")
    result = {
        "filename": " + ".join(d["filename"] for d in documents),
        "file_type": "DOCUMENT_SET",
        "vendor_name": documents[0].get("vendor_name", "Unidentified bidder"),
        "page_count": sum(d["page_count"] for d in documents),
        "evidence": {},
        "unread_pages": [],
        "documents": [
            {k: d[k] for k in ("filename", "file_type", "page_count")}
            for d in documents
        ],
    }
    for d in documents:
        for key, values in d["evidence"].items():
            result["evidence"].setdefault(key, []).extend(values)
        result["unread_pages"].extend(d["unread_pages"])
    for key, values in result["evidence"].items():
        unique = list(dict.fromkeys(e["value"] for e in values))
        # Ambiguous attachment identities/claims must be resolved by an officer.
        if len(unique) == 1:
            result[key] = unique[0]
    result["all_pans"] = list(
        dict.fromkeys(e["value"] for e in result["evidence"].get("pan", []))
    )
    result["all_gstins"] = list(
        dict.fromkeys(e["value"] for e in result["evidence"].get("gstin", []))
    )
    result["msme_claimed"] = any(d.get("msme_claimed") for d in documents)
    result["is_msme"] = False
    result["raw_text"] = "\n".join(d["raw_text"] for d in documents)
    return result


def extract_tender_rfp_data(path):
    pages = extract_pages(path)
    text = "\n".join(p["text"] for p in pages)
    patterns = {
        "budget_inr": r"(?:Estimated\s+(?:Tender\s+)?Value|Budget)\s*[:=]\s*(?:INR|Rs\.?|₹)?\s*([\d,]+(?:\.\d+)?)\s*(crores?|cr\b|lakhs?)?",
        "emd_required_inr": r"(?:EMD|Earnest Money Deposit)\s*[:=]\s*(?:INR|Rs\.?|₹)?\s*([\d,]+(?:\.\d+)?)\s*(crores?|cr\b|lakhs?)?",
        "min_turnover_cr": r"(?:Minimum\s+(?:Annual\s+)?Turnover|Average Annual Turnover)\s*[:=]\s*(?:INR|Rs\.?|₹)?\s*([\d,]+(?:\.\d+)?)\s*(crores?|cr\b|lakhs?)?",
    }
    result = {
        key: None for key in [*patterns, "min_local_content_pct", "min_warranty_years"]
    }
    for key, pattern in patterns.items():
        m = re.search(pattern, text, re.I)
        if m:
            result[key] = money(m.group(1), m.group(2)) / (
                10000000 if key == "min_turnover_cr" else 1
            )
    m = re.search(
        r"(?:Minimum\s+)?Local Content\s*[:=]\s*(\d+(?:\.\d+)?)\s*%", text, re.I
    )
    if m and float(m.group(1)) <= 100:
        result["min_local_content_pct"] = float(m.group(1))
    m = re.search(
        r"(?:Minimum\s+)?Warranty\s*[:=]\s*(\d+)\s*[- ]?\s*years?", text, re.I
    )
    if m:
        result["min_warranty_years"] = float(m.group(1))
    m = re.search(r"GEM/\d{4}/[A-Z]/\d+", text, re.I)
    return {
        "tender_id": m.group(0) if m else None,
        "title": Path(path).stem,
        "filename": Path(path).name,
        "requirements": result,
        "unread_pages": [p["page"] for p in pages if p["error"]],
        "needs_officer_confirmation": True,
    }


extract_pdf_data = extract_document_data
